from pathlib import Path
import re,zipfile
from PIL import Image
from pptx import Presentation
root=Path('D:/mcp/arts/American Gothic');base='Grant_Wood_Three_Works'
for lang in ('tw','en'):
 d=root/('tw_slides' if lang=='tw' else 'en_slides')
 p=Presentation(d/f'{base}.pptx')
 ls=(d/f'{base}_subtitle.txt').read_text(encoding='utf-8').splitlines()
 print(lang,'slides',len(p.slides),'size',p.slide_width,p.slide_height,'lines',len(ls))
 for i in range(0,len(ls),2):
  text=ls[i+1]
  print(i//2+1,len(re.findall(r"[A-Za-z]+(?:'[A-Za-z]+)?",text)) if lang=='en' else len(re.sub(r'\s','',text)))
 print('end shapes',len(p.slides[-1].shapes),'layout',p.slides[-1].slide_layout.name)
for p in sorted((root/'slide_backgrounds').glob(base+'_background_*.png')):
 print(p.name,Image.open(p).size)
