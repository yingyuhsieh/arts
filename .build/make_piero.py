from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image, ImageDraw
import shutil

root=Path(r"D:\mcp\arts\Piero della Francesca")
img=root/"images"; twdir=root/"tw_slides"; endir=root/"en_slides"; bgd=root/"slide_backgrounds"
for d in (twdir,endir,bgd): d.mkdir(exist_ok=True)
bgsrc=Path(r"C:\Users\yingyuhsieh\.codex\generated_images\01a0fa0e-cde3-7451-80d2-7e472e3d9272\exec-9d8d6a49-168a-4fa7-9917-256e17420c7e.png")
if bgsrc.exists():
    for i in range(1,11): shutil.copy2(bgsrc,bgd/f"piero_della_francesca_background_{i:03}.png")
else:
    for i in range(1,11):
        im=Image.new("RGB",(1280,720),(242,238,228)); ImageDraw.Draw(im).rectangle((0,650,1280,720),fill=(198,139,119)); im.save(bgd/f"piero_della_francesca_background_{i:03}.png")
data=[
("皮耶羅・德拉・弗朗切斯卡","秩序、光與文藝復興的靜謐","Piero della Francesca","The quiet order of the Early Renaissance","IMG-001_The_Resurrection.jpg"),
("數學讓神聖變得可見","透視、比例與寧靜的光","When mathematics becomes visible","Perspective, proportion, and calm light","IMG-002_The_Baptism_of_Christ.jpg"),
("慈悲聖母多聯畫","斗篷成為一座庇護的建築","The Polyptych of Mercy","A cloak becomes an architecture of shelter","IMG-003_Polyptych_of_the_Misericordia.jpg"),
("真十字架傳奇","一整座教堂裡的歷史與夢","The Legend of the True Cross","History and dream across an entire chapel","IMG-005_Legend_of_the_True_Cross_representative_scene.jpg"),
("君士坦丁之夢","夜景中的光，像一個信仰的瞬間","The Dream of Constantine","Light at night, held like a moment of faith","IMG-006_Dream_of_Constantine.jpg"),
("基督鞭笞","兩個空間，一個持續的謎","The Flagellation of Christ","Two spaces, one enduring mystery","IMG-008_The_Flagellation_of_Christ.jpg"),
("烏爾比諾公爵夫婦","側面肖像與無限遠的翁布里亞","The Dukes of Urbino","Profile portraits and an Umbrian horizon","IMG-010A_Diptych_fronts.jpg"),
("布雷拉祭壇畫","建築、家族與沉睡的聖嬰","The Brera Altarpiece","Architecture, family, and the sleeping child","IMG-011_Montefeltro_Altarpiece_Brera_Madonna.jpg"),
("一種永恆的秩序","皮耶羅讓觀看變成冥想","An order that feels eternal","Piero turns looking into contemplation","IMG-013_The_Nativity.jpg"),
("感謝觀看","Our Famous Artists","Thank you for watching","Our Famous Artists",None)]
tw=["在早期文藝復興的畫家之中，皮耶羅・德拉・弗朗切斯卡的作品最容易讓人安靜下來。比例、空間與色彩彼此對齊，像一座慢慢展開的建築。","皮耶羅不只是畫家，也是數學家與透視理論的作者。《基督受洗》中，中央的基督、白鴿與水碗幾乎排成一條垂直軸線。","《慈悲聖母多聯畫》把保護變成可見的空間。聖母張開斗篷，將跪拜者收納在兩側，褶皺像石造拱頂一樣穩定。","《真十字架傳奇》把每個事件安置在清晰的建築和風景中，人物像舞台上的雕塑，歷史因此同時顯得遙遠又在眼前。","《君士坦丁之夢》以黑暗包圍帳篷與沉睡的皇帝，真正的焦點來自左側的天使光芒。光不是裝飾，而是事件本身。","《基督鞭笞》以中央柱子分隔兩個空間。左側依嚴格透視法向後退，右側三人卻像停留在另一個時間層。","烏菲茲的公爵夫婦雙聯肖像，把兩張嚴格的側面臉孔放在翁布里亞遠景前，背面的凱旋場景又把私人肖像轉成家族記憶。","《布雷拉祭壇畫》讓古典建築成為畫面的骨架，鴕鳥蛋的象徵意義則留在精確空間裡供人思考。","皮耶羅改變了我們觀看故事的方式。人物像沉靜的柱子，風景像可以測量的距離，光把不同部分連成整體。","感謝你的觀看，歡迎回到 Our Famous Artists，我們下次見。"]
en=["Among Early Renaissance painters, Piero della Francesca makes a viewer slow down. Proportion, space, and color settle into alignment, like an architecture unfolding at a measured pace.","Piero was also a mathematician and an author on perspective. In The Baptism of Christ, Christ, the dove, and John’s bowl form an almost exact vertical axis.","The Polyptych of Mercy turns protection into visible space. The Virgin gathers the kneeling donors beneath her cloak, whose folds hold their shape like stone.","The Legend of the True Cross fills the chapel at San Francesco in Arezzo with a vast painted history. Piero places each event inside lucid architecture and landscape.","The Dream of Constantine is a memorable night scene. The tent and sleeping emperor are held in darkness, while an angel’s light enters from the left.","The Flagellation of Christ places a column between two spaces. Precision and uncertainty remain together.","The Uffizi diptych places two severe profile portraits against an Umbrian horizon. On the reverse, a triumphal scene turns private portraiture into family memory.","The Brera Altarpiece makes classical architecture the skeleton of the image. Its meanings remain open inside the exact space.","Piero changed how stories could be seen. Figures become quiet columns, landscapes become measurable distance, and light binds the parts into one whole.","Thank you for watching. Return to Our Famous Artists for another journey through art."]
def build(out,lang,narr):
    r=Presentation(); r.slide_width=Inches(13.333); r.slide_height=Inches(7.5)
    for i,row in enumerate(data):
        s=r.slides.add_slide(r.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor(242,238,228)
        if i<9:
            s.shapes.add_picture(str(img/row[4]), Inches(6.6), Inches(1.1), width=Inches(5.8), height=Inches(5.2))
            t=row[0] if lang=="tw" else row[2]; sub=row[1] if lang=="tw" else row[3]
            box=s.shapes.add_textbox(Inches(.7),Inches(1.2),Inches(5.3),Inches(1.3)); tf=box.text_frame; tf.text=t; tf.paragraphs[0].font.size=Pt(28); tf.paragraphs[0].font.bold=True; tf.paragraphs[0].font.color.rgb=RGBColor(36,52,71)
            box=s.shapes.add_textbox(Inches(.75),Inches(2.4),Inches(5.1),Inches(1.1)); tf=box.text_frame; tf.text=sub; tf.paragraphs[0].font.size=Pt(20); tf.paragraphs[0].font.color.rgb=RGBColor(140,79,74)
        else:
            sticker=Path(r"D:\mcp\arts\.codex\skills\create-youtube-slides\assets\end_sticker.png"); s.shapes.add_picture(str(sticker), Inches(4.2), Inches(1.6), height=Inches(4.5))
    r.save(out/"piero_della_francesca.pptx")
    (out/"piero_della_francesca_subtitle.txt").write_text("\n".join(f"[Page {i+1}]\n{x}" for i,x in enumerate(narr)),encoding="utf-8")
build(twdir,"tw",tw); build(endir,"en",en)
